"""
Creates a clean, blinded human review task workbook for the 90 system replies.
All rating columns are left completely blank with data validation dropdowns
based on the LLM Judge rubric.
"""

import json
import random
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path("E:/Reply_agent")
INPUT_RATINGS_PATH = REPO_ROOT / "results" / "phase6" / "human_ratings_90.jsonl"
OUT_XLSX_PATH = REPO_ROOT / "results" / "phase6" / "reply_human_review_task.xlsx"

def create_workbook():
    with open(INPUT_RATINGS_PATH, "r", encoding="utf-8") as f:
        records = [json.loads(line) for line in f if line.strip()]
    
    assert len(records) == 90, f"Expected 90 records, got {len(records)}"

    wb = openpyxl.Workbook()
    
    # Sheet 1: Reply_Review_Task
    ws = wb.active
    ws.title = "Reply_Review_Task"
    ws.views.sheetView[0].showGridLines = True
    
    # Sheet 2: Rubric_Reference
    ref_ws = wb.create_sheet(title="Rubric_Reference")
    ref_ws.views.sheetView[0].showGridLines = True
    
    # Define Rubric in reference sheet
    rubric_rows = [
        ("Dimension", "Scale", "Definition & Criteria"),
        ("Relevance", "0", "Irrelevant: Reply does not address the customer inquiry or discusses unrelated topics."),
        ("Relevance", "1", "Partially Relevant: Touches on the customer topic but misses main problem or includes off-topic remarks."),
        ("Relevance", "2", "Fully Relevant: Directly and completely addresses the specific customer inquiry."),
        ("Grounding", "0", "Ungrounded: Makes claims contradicted by official policy or invents unverified information."),
        ("Grounding", "1", "Partially Grounded: Plausible general support advice but lacks clear policy/knowledge grounding."),
        ("Grounding", "2", "Fully Grounded: Strictly consistent with official Spotify support policy, public FAQs, or cited knowledge."),
        ("Usefulness", "0", "Not Useful: Generic evasion, circular non-help, or advice that cannot assist the customer."),
        ("Usefulness", "1", "Moderately Useful: Gives general guidance or diagnostic steps that require further customer effort."),
        ("Usefulness", "2", "Highly Useful: Clear, actionable resolution path, troubleshooting step, or appropriate specialist escalation."),
        ("Tone", "0", "Inappropriate: Rude, dismissive, robotic boilerplate, or hostile."),
        ("Tone", "1", "Acceptable: Neutral, somewhat impersonal, but polite."),
        ("Tone", "2", "Excellent: Empathetic, warm, professional, and branded customer support tone."),
        ("Critical_Error", "False", "No critical error: Compliant support behavior."),
        ("Critical_Error", "True", "Critical violation: Promises unauthorized refund, falsely claims server outage, promises backend edits, or unsafely auto-handles account security."),
    ]
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1DB954", end_color="1DB954", fill_type="solid") # Spotify green
    dark_header_fill = PatternFill(start_color="191414", end_color="191414", fill_type="solid")
    
    for r_idx, row in enumerate(rubric_rows, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ref_ws.cell(r_idx, c_idx, val)
            if r_idx == 1:
                cell.font = header_font
                cell.fill = dark_header_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ref_ws.column_dimensions["A"].width = 18
    ref_ws.column_dimensions["B"].width = 10
    ref_ws.column_dimensions["C"].width = 80

    # Main Sheet Headers
    headers = [
        "Item_ID",
        "Example_ID",
        "Blinded_System_ID",
        "Customer_Message",
        "Predicted_Reply",
        "Human_Relevance",      # 0, 1, 2
        "Human_Grounding",      # 0, 1, 2
        "Human_Usefulness",     # 0, 1, 2
        "Human_Tone",           # 0, 1, 2
        "Human_Critical_Error", # False, True
        "Annotator_ID",
        "Notes"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.font = header_font
        if col_idx <= 5:
            cell.fill = dark_header_fill # Context cols
        else:
            cell.fill = header_fill # Blank Human Input cols
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    # System blinding map: map system IDs consistently to Blinded codes
    system_blind_map = {
        "baseline_0_majority": "System_Alpha",
        "baseline_1_rules": "System_Beta",
        "main_agent_v1": "System_Gamma"
    }

    border_side = Side(border_style="thin", color="D9D9D9")
    row_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    for idx, r in enumerate(records, 1):
        row_num = idx + 1
        ws.row_dimensions[row_num].height = 65
        
        item_id = f"REV-90-{idx:03d}"
        eid = r["example_id"]
        blind_sys = system_blind_map.get(r["system_id"], "System_Unknown")
        msg = r["message"]
        reply = r["predicted_reply"]

        ws.cell(row_num, 1, item_id).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row_num, 2, eid).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row_num, 3, blind_sys).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row_num, 4, msg).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row_num, 5, reply).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Columns 6 to 12 are completely blank for human review
        for c in range(6, 13):
            cell = ws.cell(row_num, c, None)
            cell.alignment = Alignment(horizontal="center" if c < 11 else "left", vertical="top", wrap_text=True)
            cell.border = row_border
        
        for c in range(1, 6):
            ws.cell(row_num, c).border = row_border

    # Data Validations
    dv_012 = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    dv_bool = DataValidation(type="list", formula1='"False,True"', allow_blank=True)
    
    ws.add_data_validation(dv_012)
    ws.add_data_validation(dv_bool)

    dv_012.add("F2:I91")
    dv_bool.add("J2:J91")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 55
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["K"].width = 16
    ws.column_dimensions["L"].width = 30

    OUT_XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX_PATH)
    print(f"Successfully generated {OUT_XLSX_PATH} with 90 blank human review rows!")

if __name__ == "__main__":
    create_workbook()
