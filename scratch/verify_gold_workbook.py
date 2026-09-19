import openpyxl
from pathlib import Path

xlsx_path = Path("gold_human_review.xlsx")
assert xlsx_path.exists()

wb = openpyxl.load_workbook(xlsx_path)
print("Sheets in workbook:", wb.sheetnames)
assert "Gold_Human_Review" in wb.sheetnames
assert "Taxonomy_Reference" in wb.sheetnames

ws = wb["Gold_Human_Review"]
print("Max row:", ws.max_row)
print("Max column:", ws.max_column)
assert ws.max_row == 201
assert ws.max_column == 13

headers = [ws.cell(1, col).value for col in range(1, 14)]
print("Headers:", headers)

# Check first row data
row2 = [ws.cell(2, col).value for col in range(1, 14)]
print("\nRow 2 Example ID:", row2[0])
print("Row 2 Subset:", row2[1])
print("Row 2 Prior Context:", row2[2][:60] if row2[2] else "None")
print("Row 2 Target Msg:", row2[3][:60] if row2[3] else "None")
print("Row 2 Blank Human Fields (Cols 5-13):", row2[4:])
assert all(val in (None, "") for val in row2[4:]), "Human annotation fields must be blank!"

# Check last row data
row201 = [ws.cell(201, col).value for col in range(1, 14)]
print("\nRow 201 Example ID:", row201[0])
print("Row 201 Subset:", row201[1])
print("Row 201 Blank Human Fields (Cols 5-13):", row201[4:])
assert all(val in (None, "") for val in row201[4:]), "Human annotation fields must be blank!"

# Check data validation
print("\nData validations in sheet:", len(ws.data_validations.dataValidation))
for idx, dv in enumerate(ws.data_validations.dataValidation):
    print(f"  DV {idx+1}: type={dv.type}, formula={dv.formula1}, sqref={dv.sqref}")

# Check reference sheet
ref_ws = wb["Taxonomy_Reference"]
print("\nTaxonomy Reference max row:", ref_ws.max_row)
intents = [ref_ws.cell(r, 1).value for r in range(2, 10)]
print("Intents in reference sheet:", intents)

print("\nAll checks passed successfully!")
